"""Hoja de datos auxiliar (oculta) que alimenta las gráficas nativas.

Los gráficos de openpyxl leen celdas reales del libro, no listas de Python
sueltas — así que hace falta una tabla en el propio workbook por cada serie.
Se usa una hoja aparte en vez de las hojas "Semanal"/"Progresión" porque
esas están pensadas para lectura humana (más reciente arriba, filas
TOTAL+deporte mezcladas) y no dan un rango de celdas limpio para graficar.
Como esta hoja se regenera en cada sync junto con el resto del libro, los
gráficos siguen siendo 100% nativos y "vivos": si cambian los datos aguas
arriba y se regenera el Excel, el gráfico se actualiza solo.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, timedelta

from openpyxl.workbook.workbook import Workbook

from garmin_tracker.activity_model import ActivityRow
from garmin_tracker.aggregates import WeeklyAggregate, moving_average
from garmin_tracker.excel_common import HEADER_FONT, apply_number_formats, write_header
from garmin_tracker.progression import STANDARD_DISTANCES_M, AcwrPoint, PersonalBest, efficiency_factor_series

SHEET_NAME = "_ChartData"
_DURATION_FMT = "[h]:mm:ss"


@dataclass(frozen=True)
class ChartDataLayout:
    sheet_name: str
    weekly_header_row: int
    weekly_row_count: int
    ef_header_row: int
    ef_start_col: int
    ef_sports: list[str]
    ef_row_count: int
    pb_header_row: int
    pb_blocks: dict[int, tuple[int, int]]  # distancia_m -> (columna Fecha, nº de filas)
    dist_header_row: int
    dist_start_col: int
    dist_row_count: int


def write_chart_data_sheet(
    wb: Workbook,
    *,
    activity_rows: list[ActivityRow],
    continuous_totals: list[WeeklyAggregate],
    acwr_points: list[AcwrPoint],
    weekly_aggregates: list[WeeklyAggregate],
    personal_best_history: dict[int, list[PersonalBest]],
) -> ChartDataLayout:
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME)
    ws.sheet_state = "hidden"

    # --- Bloque 1 (cols A-G): volumen semanal, media móvil, carga y ACWR ---
    weekly_header_row = 1
    write_header(
        ws,
        [
            "Inicio semana", "Volumen (km)", "Media móvil 4 sem (km)",
            "Carga TRIMP (UA)", "ACWR", "ACWR mínimo", "ACWR máximo",
        ],
        row=weekly_header_row,
    )
    volumes = [w.total_distance_km for w in continuous_totals]
    ma4_volume = moving_average(volumes, window=4)
    acwr_by_week = {(a.iso_year, a.iso_week): a for a in acwr_points}

    for i, (w, ma) in enumerate(zip(continuous_totals, ma4_volume)):
        r = weekly_header_row + 1 + i
        acwr = acwr_by_week.get((w.iso_year, w.iso_week))
        ws.append(
            [
                date.fromisocalendar(w.iso_year, w.iso_week, 1),
                w.total_distance_km,
                ma,
                w.training_load,
                acwr.ratio if acwr else None,
                0.8,
                1.3,
            ]
        )
    weekly_row_count = len(continuous_totals)
    apply_number_formats(
        ws,
        ["DD/MM/YYYY", "0.00", "0.00", "0.0", "0.00", "0.00", "0.00"],
        start_row=weekly_header_row + 1,
        end_row=weekly_header_row + weekly_row_count,
    )

    # --- Bloque 2: Efficiency Factor por deporte, un deporte por columna ---
    ef_start_col = 9  # columna I, deja una separación tras el bloque 1
    ef_points = list(efficiency_factor_series(weekly_aggregates))

    sports_in_order: list[str] = []
    sport_labels: dict[str, str] = {}
    for p in ef_points:
        if p.sport_key not in sport_labels:
            sport_labels[p.sport_key] = p.sport_label
            sports_in_order.append(p.sport_key)

    weeks_with_ef = sorted({(p.iso_year, p.iso_week) for p in ef_points})
    ef_value = {((p.iso_year, p.iso_week), p.sport_key): p.value for p in ef_points}

    ef_header_row = 1
    ws.cell(row=ef_header_row, column=ef_start_col, value="Inicio semana").font = HEADER_FONT
    for col_offset, sport_key in enumerate(sports_in_order, start=1):
        ws.cell(row=ef_header_row, column=ef_start_col + col_offset, value=sport_labels[sport_key]).font = HEADER_FONT

    for i, week_key in enumerate(weeks_with_ef):
        r = ef_header_row + 1 + i
        week_start = date.fromisocalendar(week_key[0], week_key[1], 1)
        ws.cell(row=r, column=ef_start_col, value=week_start).number_format = "DD/MM/YYYY"
        for col_offset, sport_key in enumerate(sports_in_order, start=1):
            value = ef_value.get((week_key, sport_key))
            if value is not None:
                cell = ws.cell(row=r, column=ef_start_col + col_offset, value=value)
                cell.number_format = "0.000"
    ef_row_count = len(weeks_with_ef)

    # --- Bloque 3: histórico de mejores marcas (Fecha/Tiempo por distancia) ---
    pb_header_row = 1
    pb_start_col = ef_start_col + len(sports_in_order) + 2
    pb_blocks: dict[int, tuple[int, int]] = {}
    col = pb_start_col
    for distance_m in sorted(STANDARD_DISTANCES_M):
        records = personal_best_history.get(distance_m, [])
        label = STANDARD_DISTANCES_M[distance_m]
        ws.cell(row=pb_header_row, column=col, value=f"Fecha {label}").font = HEADER_FONT
        ws.cell(row=pb_header_row, column=col + 1, value=f"Tiempo {label}").font = HEADER_FONT
        for i, pb in enumerate(records):
            r = pb_header_row + 1 + i
            ws.cell(row=r, column=col, value=pb.achieved_on).number_format = "DD/MM/YYYY"
            ws.cell(row=r, column=col + 1, value=timedelta(seconds=pb.time_seconds)).number_format = _DURATION_FMT
        pb_blocks[distance_m] = (col, len(records))
        col += 3  # 2 columnas de datos + 1 de separación

    # --- Bloque 4: distribución de sesiones por deporte (para el pastel) ---
    dist_start_col = col
    dist_header_row = 1
    ws.cell(row=dist_header_row, column=dist_start_col, value="Deporte").font = HEADER_FONT
    ws.cell(row=dist_header_row, column=dist_start_col + 1, value="Sesiones").font = HEADER_FONT
    counts: dict[str, int] = {}
    for row in activity_rows:
        counts[row.sport_label] = counts.get(row.sport_label, 0) + 1
    dist_rows = sorted(counts.items(), key=lambda kv: -kv[1])
    for i, (label, count) in enumerate(dist_rows):
        r = dist_header_row + 1 + i
        ws.cell(row=r, column=dist_start_col, value=label)
        ws.cell(row=r, column=dist_start_col + 1, value=count)
    dist_row_count = len(dist_rows)

    return ChartDataLayout(
        sheet_name=SHEET_NAME,
        weekly_header_row=weekly_header_row,
        weekly_row_count=weekly_row_count,
        ef_header_row=ef_header_row,
        ef_start_col=ef_start_col,
        ef_sports=[sport_labels[s] for s in sports_in_order],
        ef_row_count=ef_row_count,
        pb_header_row=pb_header_row,
        pb_blocks=pb_blocks,
        dist_header_row=dist_header_row,
        dist_start_col=dist_start_col,
        dist_row_count=dist_row_count,
    )
