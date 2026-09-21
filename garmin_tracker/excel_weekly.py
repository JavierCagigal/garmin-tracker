"""Hoja 'Semanal': agregados por semana ISO, por deporte y total."""

from __future__ import annotations

from datetime import date, timedelta
from itertools import groupby

from openpyxl.workbook.workbook import Workbook

from garmin_tracker.aggregates import TOTAL_SPORT_KEY, WeeklyAggregate
from garmin_tracker.excel_common import FILL_TOTAL_ROW, apply_number_formats, autofit_columns, write_header

SHEET_NAME = "Semanal"

_DURATION_FMT = "[h]:mm:ss"
_COLUMNS: list[tuple[str, str | None]] = [
    ("Inicio semana (lunes)", "DD/MM/YYYY"),
    ("Año ISO", "0"),
    ("Semana ISO", "0"),
    ("Deporte", None),
    ("Sesiones", "0"),
    ("Km totales", "0.00"),
    ("Tiempo total", _DURATION_FMT),
    ("Desnivel total (m)", "0"),
    ("Ritmo medio ponderado (min/km)", "mm:ss"),
    ("FC media ponderada", "0"),
    ("Carga semanal (TRIMP, UA)", "0.0"),
    ("% tiempo suave (z1-2)", '0.0"%"'),
    ("% tiempo intenso (z3-5)", '0.0"%"'),
]


def _most_recent_week_first(weekly: list[WeeklyAggregate]) -> list[WeeklyAggregate]:
    """weekly viene en orden cronológico ascendente (TOTAL + deportes por
    semana); aquí se invierte el orden de las SEMANAS, conservando el orden
    TOTAL-luego-deportes dentro de cada una, para mostrar lo más reciente
    arriba como en la hoja Actividades."""
    week_groups = [list(group) for _key, group in groupby(weekly, key=lambda w: (w.iso_year, w.iso_week))]
    week_groups.reverse()
    return [row for group in week_groups for row in group]


def _row_values(w: WeeklyAggregate) -> list[object]:
    pace_value = (
        timedelta(minutes=w.weighted_avg_pace_min_per_km) if w.weighted_avg_pace_min_per_km is not None else None
    )
    week_start = date.fromisocalendar(w.iso_year, w.iso_week, 1)
    return [
        week_start,
        w.iso_year,
        w.iso_week,
        w.sport_label,
        w.sessions,
        w.total_distance_km,
        w.total_duration,
        w.total_elevation_gain_m,
        pace_value,
        w.weighted_avg_hr,
        w.training_load,
        w.pct_time_easy,
        w.pct_time_intense,
    ]


def write_weekly_sheet(wb: Workbook, weekly_aggregates: list[WeeklyAggregate]) -> None:
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME)

    write_header(ws, [header for header, _fmt in _COLUMNS])

    for w in _most_recent_week_first(weekly_aggregates):
        ws.append(_row_values(w))
        if w.sport_key == TOTAL_SPORT_KEY:
            for cell in ws[ws.max_row]:
                cell.fill = FILL_TOTAL_ROW

    apply_number_formats(ws, [fmt for _h, fmt in _COLUMNS], start_row=2, end_row=ws.max_row)
    autofit_columns(ws)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
