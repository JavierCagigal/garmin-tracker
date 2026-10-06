"""Hoja 'Progresión': "¿he mejorado?", separada por deporte (un bloque
debajo de otro: Correr, Natación, Ciclismo, Fuerza) y sesión a sesión en
orden cronológico, con pocas columnas para no tener que ir a la derecha y
perder la fila.

Cada sesión con ritmo se compara con la anterior del MISMO deporte y con la
media de las 4 anteriores: verde = más rápido, rojo = más lento. Caminar y
otros deportes sueltos no aparecen.
"""

from __future__ import annotations

from datetime import timedelta

from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from garmin_tracker.activity_model import ActivityRow
from garmin_tracker.excel_charts import SPORTS
from garmin_tracker.excel_common import (
    FILL_GREEN,
    FILL_RED,
    TITLE_FONT,
    apply_number_formats,
    autofit_columns,
    write_header,
)
from garmin_tracker.progression import STANDARD_DISTANCES_M, PersonalBest, pace_deltas

SHEET_NAME = "Progresión"

_DURATION_FMT = "h:mm:ss"
# Las diferencias de ritmo van en segundos: un tiempo negativo en formato
# mm:ss no se puede mostrar (sale como una fecha absurda tipo 29/12/99). Una
# sola sección, sin "+": con secciones +/- Numbers duplica el signo ("--32 s").
_SIGNED_SECONDS_FMT = '0" s"'
# En la tabla, la duración ya va en "Tiempo": no se repite como métrica.
_SKIPPED_METRICS = {"Sesiones", "Minutos"}

_PB_COLUMNS: list[tuple[str, str | None]] = [
    ("Distancia", None),
    ("Tiempo", _DURATION_FMT),
    ("Ritmo (min/km)", "mm:ss"),
    ("Fecha", "dd/mm/yyyy"),
    ("Actividad", None),
    ("activityId", "0"),
]


def _delta_fill(delta: float | None) -> object | None:
    if not delta:
        return None
    return FILL_GREEN if delta < 0 else FILL_RED  # en ritmo, menos = mejor


def _write_sport_section(ws: Worksheet, start_row: int, sport, rows: list[ActivityRow], titles: list) -> int:
    sessions = sorted((r for r in rows if r.sport_key in sport.keys), key=lambda r: (r.date, r.start_time))
    if not sessions:
        return start_row
    metrics = [m for m in sport.metrics if m.name not in _SKIPPED_METRICS]
    pace = next((m for m in metrics if m.number_format == "m:ss"), None)

    headers = ["Fecha", "Tiempo"] + [m.name for m in metrics]
    formats = ["dd/mm/yyyy", _DURATION_FMT] + [m.number_format for m in metrics]
    if pace:
        headers += ["Δ ritmo vs anterior", "Δ ritmo vs media 4 anteriores"]
        formats += [_SIGNED_SECONDS_FMT, _SIGNED_SECONDS_FMT]
    headers.append("Actividad")
    formats.append(None)

    titles.append((start_row, f"{sport.name} — sesión a sesión"))
    header_row = start_row + 1
    write_header(ws, headers, row=header_row)

    pace_values = [pace.compute([r]) if pace else None for r in sessions]
    deltas = pace_deltas([p.total_seconds() if p else None for p in pace_values])
    row = header_row + 1
    for r, delta in zip(sessions, deltas):
        values = [r.date, r.duration] + [m.compute([r]) for m in metrics]
        if pace:
            values += [
                round(delta.vs_prev) if delta.vs_prev is not None else None,
                round(delta.vs_avg4) if delta.vs_avg4 is not None else None,
            ]
            for offset, d in enumerate((delta.vs_prev, delta.vs_avg4)):
                fill = _delta_fill(d)
                if fill is not None:
                    ws.cell(row=row, column=len(metrics) + 3 + offset).fill = fill
        values.append(r.name)
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        row += 1

    apply_number_formats(ws, formats, start_row=header_row + 1, end_row=row - 1)
    return row


def _write_personal_bests_section(ws: Worksheet, start_row: int, personal_bests: dict[int, PersonalBest], titles: list) -> int:
    titles.append((start_row, "Correr — mejores marcas"))
    header_row = start_row + 1
    write_header(ws, [h for h, _fmt in _PB_COLUMNS], row=header_row)

    row = header_row + 1
    for distance_m in sorted(STANDARD_DISTANCES_M):
        pb = personal_bests.get(distance_m)
        if pb is None:
            ws.append([STANDARD_DISTANCES_M[distance_m], None, None, None, "Sin marca todavía", None])
        else:
            pace = timedelta(minutes=(pb.time_seconds / 60) / (pb.distance_m / 1000))
            ws.append(
                [
                    pb.distance_label,
                    timedelta(seconds=round(pb.time_seconds)),
                    pace,
                    pb.achieved_on,
                    pb.activity_name,
                    pb.activity_id,
                ]
            )
        row += 1

    apply_number_formats(ws, [fmt for _h, fmt in _PB_COLUMNS], start_row=header_row + 1, end_row=row - 1)
    return row


def write_progression_sheet(
    wb: Workbook,
    *,
    rows: list[ActivityRow],
    personal_bests: dict[int, PersonalBest],
) -> None:
    ws = wb.create_sheet(SHEET_NAME)
    ws.freeze_panes = "B1"  # la fecha siempre visible si hay que ir a la derecha

    titles: list[tuple[int, str]] = []
    next_row = 3
    for sport in SPORTS:
        before = next_row
        next_row = _write_sport_section(ws, next_row, sport, rows, titles)
        if sport.name == "Correr" and next_row != before:
            next_row = _write_personal_bests_section(ws, next_row + 1, personal_bests, titles)
        if next_row != before:
            next_row += 2  # dos filas en blanco entre deportes

    autofit_columns(ws)
    # Títulos y leyenda después del autoajuste, para no ensanchar la columna A.
    ws.cell(row=1, column=1, value="Verde = más rápido que antes, rojo = más lento (diferencias en segundos por km, o por 100 m en natación).")
    for row, text in titles:
        ws.cell(row=row, column=1, value=text).font = TITLE_FONT
