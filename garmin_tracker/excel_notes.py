"""Hoja "Notas": columnas manuales (RPE, sensaciones, lesiones, contexto)
ligadas por activityId.

Fusión segura: `read_existing_notes` lee el Excel generado en la ejecución
ANTERIOR (si existe) y extrae lo que el usuario haya escrito a mano, por
completo en memoria, ANTES de que `write_workbook` empiece a construir el
libro nuevo. El fichero original no se toca en ningún momento durante ese
proceso — solo se sustituye al final, de forma atómica (ver
`excel_common.save_workbook_atomic`), y solo si todo lo anterior ha
funcionado. Si el Excel previo existe pero no se puede leer como .xlsx
válido, se deja que la excepción se propague: mejor abortar la generación
completa que arriesgarse a sobrescribir notas que no hemos podido confirmar.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from garmin_tracker.activity_model import ActivityRow
from garmin_tracker.excel_common import apply_number_formats, autofit_columns, write_header

SHEET_NAME = "Notas"

_AUTO_HEADERS = ["activityId", "Fecha", "Deporte", "Nombre"]
_MANUAL_HEADERS = ["RPE percibido", "Sensaciones", "Lesiones / molestias", "Contexto"]
_HEADERS = _AUTO_HEADERS + _MANUAL_HEADERS


@dataclass(frozen=True)
class ManualNote:
    rpe: Any = None
    sensaciones: Any = None
    lesiones: Any = None
    contexto: Any = None

    def is_empty(self) -> bool:
        return all(v is None for v in (self.rpe, self.sensaciones, self.lesiones, self.contexto))


def read_existing_notes(path: Path) -> dict[int, ManualNote]:
    """Lee las notas manuales del Excel generado la vez anterior.

    Devuelve {} (nada que perder) si el fichero no existe todavía, o existe
    pero no tiene hoja "Notas" (p.ej. la primera vez que se genera con esta
    función ya presente) o su cabecera no tiene la forma esperada. Si el
    fichero existe y no se puede abrir como .xlsx válido, propaga la
    excepción a propósito.
    """
    if not path.exists():
        return {}

    wb = load_workbook(path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        return {}

    ws = wb[SHEET_NAME]
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    try:
        id_col = header.index("activityId")
        rpe_col = header.index("RPE percibido")
        sens_col = header.index("Sensaciones")
        les_col = header.index("Lesiones / molestias")
        ctx_col = header.index("Contexto")
    except ValueError:
        # Cabecera con forma inesperada (p.ej. de una versión futura/pasada
        # distinta): no arriesgamos a mapear mal las columnas.
        return {}

    notes: dict[int, ManualNote] = {}
    for row in ws.iter_rows(min_row=2):
        activity_id = row[id_col].value
        if activity_id is None:
            continue
        note = ManualNote(
            rpe=row[rpe_col].value,
            sensaciones=row[sens_col].value,
            lesiones=row[les_col].value,
            contexto=row[ctx_col].value,
        )
        if not note.is_empty():
            notes[int(activity_id)] = note
    return notes


def write_notes_sheet(
    wb: Workbook, activity_rows: list[ActivityRow], existing_notes: dict[int, ManualNote]
) -> None:
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME)
    write_header(ws, _HEADERS)

    seen_ids: set[int] = set()
    for row in activity_rows:
        note = existing_notes.get(row.activity_id, ManualNote())
        ws.append(
            [
                row.activity_id, row.date, row.sport_label, row.name,
                note.rpe, note.sensaciones, note.lesiones, note.contexto,
            ]
        )
        seen_ids.add(row.activity_id)

    # Notas de actividades que ya no aparezcan en la caché actual (no
    # debería ocurrir nunca, porque nunca borramos de la caché — pero por
    # seguridad, mejor una fila "huérfana" que perder lo escrito a mano).
    for activity_id, note in existing_notes.items():
        if activity_id in seen_ids:
            continue
        ws.append(
            [activity_id, None, None, "(ya no está en la caché)", note.rpe, note.sensaciones, note.lesiones, note.contexto]
        )

    apply_number_formats(ws, ["0", "DD/MM/YYYY", None, None, None, None, None, None], start_row=2, end_row=ws.max_row)
    autofit_columns(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
