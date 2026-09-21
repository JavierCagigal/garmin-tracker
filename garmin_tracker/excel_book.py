"""Generación del libro Excel (.xlsx) a partir de filas ya calculadas.

Compatible con LibreOffice Calc: se usa openpyxl puro (formato OOXML
estándar), sin ninguna característica exclusiva de Microsoft Excel.

Escritura atómica: se escribe primero a un fichero temporal en el mismo
directorio y solo se sustituye el definitivo si `wb.save()` termina sin
errores, para no dejar nunca un .xlsx corrupto o a medias si algo falla a
mitad de la generación.
"""

from __future__ import annotations

import os
import tempfile
from datetime import timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook as WorkbookType
from openpyxl.worksheet.worksheet import Worksheet

from garmin_tracker.activity_model import ActivityRow

ACTIVITIES_SHEET_NAME = "Actividades"

# (cabecera, ancho de columna, formato numérico o None para texto)
_DURATION_FMT = "[h]:mm:ss"
_COLUMNS: list[tuple[str, int, str | None]] = [
    ("Fecha", 12, "DD/MM/YYYY"),
    ("Hora inicio", 11, "HH:MM:SS"),
    ("Deporte", 20, None),
    ("Nombre", 32, None),
    ("Distancia (km)", 14, "0.00"),
    ("Tiempo total", 12, _DURATION_FMT),
    ("Tiempo en movimiento", 19, _DURATION_FMT),
    ("Ritmo medio (min/km)", 18, "mm:ss"),
    ("Velocidad media (km/h)", 19, "0.0"),
    ("FC media", 10, "0"),
    ("FC máx", 10, "0"),
    ("Desnivel + (m)", 14, "0"),
    ("Calorías", 10, "0"),
    ("Cadencia media", 14, "0.0"),
    ("Potencia media (W)", 16, "0"),
    ("TE aeróbico", 12, "0.0"),
    ("TE anaeróbico", 13, "0.0"),
    ("VO2max", 10, "0.0"),
    ("FC Zona 1", 11, _DURATION_FMT),
    ("FC Zona 2", 11, _DURATION_FMT),
    ("FC Zona 3", 11, _DURATION_FMT),
    ("FC Zona 4", 11, _DURATION_FMT),
    ("FC Zona 5", 11, _DURATION_FMT),
    ("Temp mín (°C)", 13, "0.0"),
    ("Temp máx (°C)", 13, "0.0"),
    ("activityId", 16, "0"),
]


def _row_values(row: ActivityRow) -> list[object]:
    pace_value = timedelta(minutes=row.avg_pace_min_per_km) if row.avg_pace_min_per_km is not None else None
    zone_values = [timedelta(seconds=z) if z is not None else None for z in row.hr_zone_seconds]

    return [
        row.date,
        row.start_time,
        row.sport_label,
        row.name,
        row.distance_km,
        row.duration,
        row.moving_duration,
        pace_value,
        row.avg_speed_kmh,
        row.avg_hr,
        row.max_hr,
        row.elevation_gain_m,
        row.calories,
        row.avg_cadence,
        row.avg_power,
        row.aerobic_te,
        row.anaerobic_te,
        row.vo2max,
        *zone_values,
        row.temp_min_c,
        row.temp_max_c,
        row.activity_id,
    ]


def write_activities_sheet(wb: WorkbookType, rows: list[ActivityRow]) -> Worksheet:
    if ACTIVITIES_SHEET_NAME in wb.sheetnames:
        del wb[ACTIVITIES_SHEET_NAME]
    ws = wb.create_sheet(ACTIVITIES_SHEET_NAME, 0)

    headers = [header for header, _width, _fmt in _COLUMNS]
    ws.append(headers)
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font

    for row in rows:
        ws.append(_row_values(row))

    for col_index, (_header, width, number_format) in enumerate(_COLUMNS, start=1):
        letter = get_column_letter(col_index)
        ws.column_dimensions[letter].width = width
        if number_format:
            for data_row in range(2, ws.max_row + 1):
                ws.cell(row=data_row, column=col_index).number_format = number_format

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def write_workbook(rows: list[ActivityRow], path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)  # Workbook() crea una hoja "Sheet" vacía por defecto

    write_activities_sheet(wb, rows)

    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=".tmp-", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, path)
    except Exception:
        tmp_path.unlink(missing_ok=True)
        raise
