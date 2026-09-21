"""Utilidades compartidas por los distintos escritores de hojas Excel.

Compatible con LibreOffice Calc: solo openpyxl puro (OOXML estándar), nada
exclusivo de Microsoft Excel.
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FONT = Font(bold=True)
TITLE_FONT = Font(bold=True, size=13)

# Verde = mejora, rojo = empeora, gris = muestra demasiado pequeña para concluir nada.
FILL_GREEN = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
FILL_RED = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
FILL_GRAY = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
FILL_TOTAL_ROW = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")


def write_header(ws: Worksheet, headers: list[str], row: int = 1) -> None:
    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_index, value=header)
        cell.font = HEADER_FONT


def apply_number_formats(
    ws: Worksheet,
    number_formats: list[str | None],
    start_row: int,
    end_row: int,
    col_offset: int = 0,
) -> None:
    for i, number_format in enumerate(number_formats, start=1):
        if not number_format:
            continue
        col_index = col_offset + i
        for data_row in range(start_row, end_row + 1):
            ws.cell(row=data_row, column=col_index).number_format = number_format


def autofit_columns(ws: Worksheet, min_width: int = 9, max_width: int = 42, padding: int = 2) -> None:
    """Aproximación de autoajuste de ancho de columna (openpyxl no ofrece
    autofit real): mide la longitud del texto visible de cada celda."""
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = len(str(cell.value))
            if length > widths.get(cell.column, 0):
                widths[cell.column] = length
    for col_index, content_width in widths.items():
        letter = get_column_letter(col_index)
        ws.column_dimensions[letter].width = max(min_width, min(max_width, content_width + padding))


def save_workbook_atomic(wb: Workbook, path: Path) -> None:
    """Escribe primero a un fichero temporal en el mismo directorio y solo
    sustituye el definitivo si `wb.save()` termina sin errores, para no
    dejar nunca un .xlsx corrupto o a medias si algo falla a mitad de la
    generación."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=".tmp-", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, path)
    except Exception:
        tmp_path.unlink(missing_ok=True)
        raise
