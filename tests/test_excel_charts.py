from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from datetime import date, time

from openpyxl.chart import BarChart, LineChart

from garmin_tracker.activity_model import build_activity_row
from garmin_tracker.aggregates import aggregate_by_week
from garmin_tracker.excel_book import write_workbook
from garmin_tracker.progression import current_personal_bests
from tests.fixtures import activity


def _build_multi_week_rows() -> list:
    raws = [
        activity(1, "2026-08-03 08:00:00", sport="running", distance_m=5000, duration_s=1500),
        activity(2, "2026-08-04 08:00:00", sport="lap_swimming", distance_m=400, duration_s=480),
        activity(3, "2026-08-06 08:00:00", sport="lap_swimming", distance_m=600, duration_s=840),
        # semana del 10/08: descanso total
        activity(4, "2026-08-17 08:00:00", sport="cycling", distance_m=20000, duration_s=3600),
        activity(5, "2026-08-18 08:00:00", sport="treadmill_running", distance_m=10000, duration_s=3000),
    ]
    return [build_activity_row(r) for r in raws]


def _write(rows, out_path: Path) -> None:
    weekly_aggregates = aggregate_by_week(rows)
    write_workbook(
        out_path,
        activity_rows=rows,
        weekly_aggregates=weekly_aggregates,
        personal_bests=current_personal_bests(rows),
    )


def test_one_sheet_per_sport_with_sport_named_charts(tmp_path: Path) -> None:
    out_path = tmp_path / "entrenamientos.xlsx"
    _write(_build_multi_week_rows(), out_path)

    wb = load_workbook(out_path)
    assert "Fuerza" not in wb.sheetnames  # sin sesiones de fuerza, sin pestaña
    for name in ("Correr", "Natación", "Ciclismo"):
        ws = wb[name]
        assert ws._images == []  # gráficos nativos, no imágenes
        assert ws._charts
        for chart in ws._charts:
            title = "".join(r.t for p in chart.title.tx.rich.p for r in (p.r or []))
            assert title.startswith(f"{name} · ")  # nunca una "FC media" sin deporte
            if isinstance(chart, BarChart):
                assert chart.gapWidth == 30
    assert any(isinstance(c, LineChart) for c in wb["Correr"]._charts)


def test_sport_table_is_weekly_continuous_with_weighted_pace(tmp_path: Path) -> None:
    out_path = tmp_path / "entrenamientos.xlsx"
    _write(_build_multi_week_rows(), out_path)

    ws = load_workbook(out_path)["Natación"]
    weeks = [ws.cell(row=r, column=1).value.date() for r in range(4, ws.max_row + 1)]
    assert weeks == [date(2026, 8, 3), date(2026, 8, 10), date(2026, 8, 17)]  # incluye descansos

    # semana del 3 ago: 2 sesiones, 1000 m en 1320 s -> 2:12 /100 m
    # (ponderado, no la media de 2:00 y 2:20)
    assert [ws.cell(row=4, column=c).value for c in (2, 3, 4)] == [2, 1000, time(0, 2, 12)]
    assert ws["C5"].value is None  # semana sin natación: fila vacía, sin barra "0"

    # la cinta cuenta como "Correr": 5 km semana 1 + 10 km semana 3
    run = load_workbook(out_path)["Correr"]
    assert [run.cell(row=r, column=3).value for r in (4, 5, 6)] == [5, None, 10]


def test_no_sport_sheets_when_no_activities(tmp_path: Path) -> None:
    out_path = tmp_path / "entrenamientos.xlsx"
    _write([], out_path)

    wb = load_workbook(out_path)
    assert not set(wb.sheetnames) & {"Correr", "Natación", "Ciclismo", "Fuerza"}
