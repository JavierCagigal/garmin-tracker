from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from garmin_tracker.activity_model import build_activity_row
from garmin_tracker.aggregates import aggregate_by_week
from garmin_tracker.excel_book import write_workbook
from garmin_tracker.excel_notes import SHEET_NAME as NOTES_SHEET_NAME
from garmin_tracker.excel_notes import ManualNote, read_existing_notes
from garmin_tracker.progression import current_personal_bests
from tests.fixtures import running_activity, strength_activity


def _write(rows, out_path: Path) -> None:
    weekly_aggregates = aggregate_by_week(rows)
    write_workbook(
        out_path,
        activity_rows=rows,
        weekly_aggregates=weekly_aggregates,
        personal_bests=current_personal_bests(rows),
    )


def test_read_existing_notes_when_file_does_not_exist_yet(tmp_path: Path) -> None:
    assert read_existing_notes(tmp_path / "no_existe.xlsx") == {}


def test_read_existing_notes_when_file_has_no_notes_sheet(tmp_path: Path) -> None:
    rows = [build_activity_row(running_activity())]
    out_path = tmp_path / "sin_notas.xlsx"
    _write(rows, out_path)

    # Se borra la hoja Notas a mano, para simular un Excel de una versión
    # anterior a esta funcionalidad.
    wb = load_workbook(out_path)
    del wb[NOTES_SHEET_NAME]
    wb.save(out_path)

    assert read_existing_notes(out_path) == {}


def test_manual_notes_survive_regeneration(tmp_path: Path) -> None:
    """El caso crítico: el usuario escribe a mano en 'Notas', se vuelve a
    sincronizar y regenerar el Excel (con una actividad nueva de por medio),
    y lo escrito a mano NO debe perderse."""
    activity_a = build_activity_row(running_activity(activity_id=1001))
    out_path = tmp_path / "entrenamientos.xlsx"

    _write([activity_a], out_path)

    # El usuario abre el Excel y rellena la fila de la actividad 1001 a mano.
    wb = load_workbook(out_path)
    ws = wb[NOTES_SHEET_NAME]
    header = [c.value for c in ws[1]]
    row_idx = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == 1001)
    ws.cell(row=row_idx, column=header.index("RPE percibido") + 1, value=7)
    ws.cell(row=row_idx, column=header.index("Sensaciones") + 1, value="Piernas cargadas")
    ws.cell(row=row_idx, column=header.index("Lesiones / molestias") + 1, value="Molestia rodilla derecha")
    ws.cell(row=row_idx, column=header.index("Contexto") + 1, value="Dormí mal")
    wb.save(out_path)

    # Llega una actividad nueva y se regenera el Excel (como haría un sync real).
    activity_b = build_activity_row(strength_activity(activity_id=1002))
    _write([activity_b, activity_a], out_path)

    wb2 = load_workbook(out_path)
    ws2 = wb2[NOTES_SHEET_NAME]
    header2 = [c.value for c in ws2[1]]
    row_idx2 = next(r for r in range(2, ws2.max_row + 1) if ws2.cell(row=r, column=1).value == 1001)

    assert ws2.cell(row=row_idx2, column=header2.index("RPE percibido") + 1).value == 7
    assert ws2.cell(row=row_idx2, column=header2.index("Sensaciones") + 1).value == "Piernas cargadas"
    assert ws2.cell(row=row_idx2, column=header2.index("Lesiones / molestias") + 1).value == "Molestia rodilla derecha"
    assert ws2.cell(row=row_idx2, column=header2.index("Contexto") + 1).value == "Dormí mal"

    # La actividad nueva aparece con las columnas manuales vacías, sin reventar nada.
    row_idx_b = next(r for r in range(2, ws2.max_row + 1) if ws2.cell(row=r, column=1).value == 1002)
    assert ws2.cell(row=row_idx_b, column=header2.index("RPE percibido") + 1).value is None


def test_read_existing_notes_ignores_empty_rows(tmp_path: Path) -> None:
    rows = [build_activity_row(running_activity(activity_id=2001))]
    out_path = tmp_path / "entrenamientos.xlsx"
    _write(rows, out_path)  # sin notas manuales todavía

    notes = read_existing_notes(out_path)
    assert notes == {}  # nada escrito a mano -> nada que fusionar


def test_read_existing_notes_raises_on_corrupt_file(tmp_path: Path) -> None:
    bad_file = tmp_path / "corrupto.xlsx"
    bad_file.write_bytes(b"esto no es un xlsx valido")

    with pytest.raises(Exception):
        read_existing_notes(bad_file)


def test_manual_note_is_empty() -> None:
    assert ManualNote().is_empty() is True
    assert ManualNote(rpe=5).is_empty() is False
