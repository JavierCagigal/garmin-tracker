from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from garmin_tracker.activity_model import build_activity_row
from garmin_tracker.aggregates import aggregate_by_week
from garmin_tracker.excel_book import ACTIVITIES_SHEET_NAME, write_workbook
from garmin_tracker.excel_common import FILL_GREEN
from garmin_tracker.excel_notes import SHEET_NAME as NOTES_SHEET_NAME
from garmin_tracker.excel_progression import SHEET_NAME as PROGRESSION_SHEET_NAME
from garmin_tracker.excel_weekly import SHEET_NAME as WEEKLY_SHEET_NAME
from garmin_tracker.progression import current_personal_bests
from tests.fixtures import activity, indoor_cycling_activity, running_activity, strength_activity


def _write(rows, out_path: Path) -> None:
    """Reproduce el mismo pipeline que scripts/generate_excel.py, para no
    duplicar la orquestación en cada test."""
    weekly_aggregates = aggregate_by_week(rows)
    write_workbook(
        out_path,
        activity_rows=rows,
        weekly_aggregates=weekly_aggregates,
        personal_bests=current_personal_bests(rows),
    )


def test_write_workbook_creates_all_expected_sheets(tmp_path: Path) -> None:
    rows = [
        build_activity_row(running_activity()),
        build_activity_row(indoor_cycling_activity()),
        build_activity_row(strength_activity()),
    ]
    out_path = tmp_path / "entrenamientos.xlsx"

    _write(rows, out_path)

    assert out_path.exists()
    wb = load_workbook(out_path)
    assert wb.sheetnames == [
        ACTIVITIES_SHEET_NAME,
        WEEKLY_SHEET_NAME,
        PROGRESSION_SHEET_NAME,
        "Correr",
        "Ciclismo",
        "Fuerza",  # sin natación en los datos: no hay pestaña vacía
        NOTES_SHEET_NAME,
    ]


def test_activities_sheet_structure(tmp_path: Path) -> None:
    rows = [
        build_activity_row(running_activity()),
        build_activity_row(indoor_cycling_activity()),
        build_activity_row(strength_activity()),
    ]
    out_path = tmp_path / "entrenamientos.xlsx"
    _write(rows, out_path)

    ws = load_workbook(out_path)[ACTIVITIES_SHEET_NAME]
    assert ws["A1"].value == "Fecha"
    assert ws["D1"].value == "Nombre"
    assert ws.cell(row=1, column=ws.max_column).value == "activityId"
    assert ws["A1"].font.bold is True
    assert ws.max_row == 4  # 3 filas de datos + cabecera

    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == ws.dimensions

    date_col, duration_col = 1, 6
    assert ws.cell(row=2, column=date_col).number_format == "dd/mm/yyyy"
    assert ws.cell(row=2, column=duration_col).number_format == "h:mm:ss"


def test_weekly_sheet_has_total_and_sport_rows(tmp_path: Path) -> None:
    rows = [
        build_activity_row(running_activity()),
        build_activity_row(strength_activity()),
    ]
    out_path = tmp_path / "entrenamientos.xlsx"
    _write(rows, out_path)

    ws = load_workbook(out_path)[WEEKLY_SHEET_NAME]
    assert ws["A1"].value == "Inicio semana (lunes)"
    sport_column_values = [ws.cell(row=r, column=4).value for r in range(2, ws.max_row + 1)]
    assert "Total" in sport_column_values
    assert ws.freeze_panes == "A2"


def test_progression_sheet_is_split_by_sport_session_by_session(tmp_path: Path) -> None:
    rows = [
        build_activity_row(activity(1, "2026-09-01 08:00:00", sport="running", distance_m=5000, duration_s=1500)),
        build_activity_row(activity(2, "2026-09-03 08:00:00", sport="treadmill_running", distance_m=5000, duration_s=1450)),
        build_activity_row(activity(3, "2026-09-02 08:00:00", sport="walking", distance_m=3000, duration_s=1800)),
        build_activity_row(strength_activity()),
    ]
    out_path = tmp_path / "entrenamientos.xlsx"
    _write(rows, out_path)

    ws = load_workbook(out_path)[PROGRESSION_SHEET_NAME]
    col_a = [c.value for c in ws["A"]]
    assert "Correr — sesión a sesión" in col_a
    assert "Correr — mejores marcas" in col_a
    assert "Fuerza — sesión a sesión" in col_a
    assert not any("Caminar" in str(v) for row in ws.iter_rows(values_only=True) for v in row)

    header_row = col_a.index("Correr — sesión a sesión") + 2
    headers = [c.value for c in ws[header_row]]
    assert len([h for h in headers if h]) <= 8  # sin tener que ir a la derecha
    # 2ª sesión (la cinta cuenta como correr): 4:50 vs 5:00 -> -10 s, en verde
    delta = ws.cell(row=header_row + 2, column=headers.index("Δ ritmo vs anterior") + 1)
    assert delta.value == -10
    assert delta.fill.fgColor.rgb == FILL_GREEN.fgColor.rgb


def test_write_workbook_is_idempotent_when_regenerated(tmp_path: Path) -> None:
    rows = [build_activity_row(running_activity())]
    out_path = tmp_path / "entrenamientos.xlsx"

    _write(rows, out_path)
    _write(rows, out_path)  # segunda pasada no debe duplicar hojas ni romper nada

    wb = load_workbook(out_path)
    assert wb.sheetnames == [
        ACTIVITIES_SHEET_NAME,
        WEEKLY_SHEET_NAME,
        PROGRESSION_SHEET_NAME,
        "Correr",
        NOTES_SHEET_NAME,
    ]
    assert wb[ACTIVITIES_SHEET_NAME].max_row == 2
